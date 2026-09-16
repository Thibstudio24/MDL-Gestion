"""Exports CSV / XLSX / PDF produits en mémoire (BytesIO), sans dépendance système."""
from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core.models import Setting

MONEY_FORMAT = '#,##0.00\\ "€"'


def brand_lines() -> tuple[str, str]:
    brand = Setting.brand()
    nom = brand.get("nom") or "MDL"
    lycee = brand.get("lycee") or ""
    return nom, ("%s — %s" % (nom, lycee)) if lycee else nom


def _csv_response(filename: str, headers, rows, delimiter: str = ";") -> HttpResponse:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\r\n")
    if headers:
        writer.writerow(headers)
    for row in rows:
        writer.writerow([_flat(cell) for cell in row])
    payload = "\ufeff" + buffer.getvalue()  # BOM UTF-8 pour Excel
    response = HttpResponse(payload.encode("utf-8"), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="%s"' % filename
    return response


def _flat(value):
    if isinstance(value, Decimal):
        return "{:.2f}".format(value).replace(".", ",")
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    return value


def csv_response(filename: str, headers, rows, delimiter: str = ";") -> HttpResponse:
    return _csv_response(filename, headers, rows, delimiter)


def xlsx_response(filename: str, sheets, freeze: str = "A2", autofilter: bool = True) -> HttpResponse:
    """sheets = [(titre, [en-têtes], [lignes], {colonne: largeur})]"""
    book = Workbook()
    book.remove(book.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="33556E")
    thin = Side(style="thin", color="D3DBE3")
    for title, headers, rows, widths in sheets:
        sheet = book.create_sheet(title[:31] or "Feuille")
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            sheet.append([_cell(value) for value in row])
        for index, column in enumerate(headers, start=1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = (widths or {}).get(column, max(12, min(38, len(str(column)) + 6)))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = Border(bottom=thin)
                if isinstance(cell.value, (int, float, Decimal)):
                    cell.number_format = MONEY_FORMAT
        sheet.freeze_panes = freeze
        if autofilter and sheet.max_row > 1:
            sheet.auto_filter.ref = "A1:%s%d" % (get_column_letter(max(1, len(headers))), sheet.max_row)
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="%s"' % filename
    return response


def _cell(value):
    if isinstance(value, (datetime,)):
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    return value


def pdf_landscape_response(
    filename: str,
    title: str,
    headers,
    rows,
    subtitle: str = "",
    totals=None,
    column_widths=None,
    footer_note: str = "",
) -> HttpResponse:
    """PDF A4 paysage avec en-tête, pied de page et ligne de totaux."""
    buffer = io.BytesIO()
    nom, full = brand_lines()
    generated = timezone.localtime().strftime("%d/%m/%Y %H:%M")

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#5a6570"))
        canvas.drawString(15 * mm, A4[1] - 10 * mm, full)
        canvas.drawRightString(landscape(A4)[0] - 15 * mm, A4[1] - 10 * mm, "%s — généré le %s" % (title, generated))
        canvas.drawString(15 * mm, 10 * mm, footer_note or nom)
        canvas.drawRightString(landscape(A4)[0] - 15 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.setStrokeColor(colors.HexColor("#d3dbe3"))
        canvas.line(15 * mm, 13 * mm, landscape(A4)[0] - 15 * mm, 13 * mm)
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=18 * mm,
        bottomMargin=16 * mm,
        title=title,
        author=nom,
    )
    styles = getSampleStyleSheet()
    heading = ParagraphStyle("h", parent=styles["Title"], fontSize=15, spaceAfter=4)
    sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#5a6570"))
    story = [Paragraph(title, heading)]
    if subtitle:
        story.append(Paragraph(subtitle, sub))
    story.append(Spacer(1, 6))
    data = [[str(h) for h in headers]] + [[_pdf_cell(c) for c in row] for row in rows]
    if totals:
        data.append([_pdf_cell(c) for c in totals])
    table = Table(data, colWidths=column_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#33556E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ALIGN", (0, 0), (-1, 0), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F8")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D3DBE3")),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]
    if totals:
        style.append(("FONTNAME", (0, len(data) - 1), (-1, len(data) - 1), "Helvetica-Bold"))
        style.append(("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#E8EDF2")))
    table.setStyle(TableStyle(style))
    story.append(table)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="%s"' % filename
    return response


def _pdf_cell(value):
    if isinstance(value, Decimal):
        return "{:,.2f}".format(value).replace(",", " ").replace(".", ",") + " €"
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if value is None:
        return ""
    return str(value)


def json_download(filename: str, payload) -> HttpResponse:
    body = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    response = HttpResponse(body.encode("utf-8"), content_type="application/json; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="%s"' % filename
    return response


def mask_secret(value: str) -> str:
    return "***" if value else ""


def user_export(user) -> dict:
    """Export JSON des données d'un compte (RGPD)."""
    payload = {
        "export": "MDL Gestion",
        "genere_le": timezone.localtime().isoformat(),
        "compte": {
            "email": user.email,
            "prenom": user.first_name,
            "nom": user.last_name,
            "fonction": getattr(user, "display_function", ""),
            "telephone": getattr(user, "phone", ""),
            "mention": getattr(user, "gender", ""),
            "interne": getattr(user, "is_boarder", False),
            "statut": user.status,
            "role": getattr(user.role, "name", None) if getattr(user, "role", None) else None,
            "derniere_activite": user.last_seen.isoformat() if user.last_seen else None,
            "preferences": getattr(user, "prefs", {}),
            "a2f_active": bool(getattr(user, "totp_enabled", False)),
        },
        "historique_roles": [
            {"role": rm.role.name, "depuis": rm.since.isoformat(), "jusquau": rm.until.isoformat() if rm.until else None}
            for rm in user.role_memberships.select_related("role").all()
        ],
        "textes_acceptes": [
            {"document": la.document.title, "version": la.version, "le": la.accepted_at.isoformat(), "ip": la.ip}
            for la in user.legal_acceptances.select_related("document").all()
        ],
    }
    for label, loader in (
        ("ecritures", lambda: _finance_entries(user)),
        ("documents", lambda: _documents(user)),
        ("disponibilites", lambda: _availabilities(user)),
        ("menages", lambda: _assignments(user)),
        ("messages", lambda: _messages(user)),
        ("connexions", lambda: _logins(user)),
    ):
        try:
            payload[label] = loader()
        except Exception:
            payload[label] = []
    return payload


def _finance_entries(user):
    from finance.models import Entry

    return [
        {"date": e.day.isoformat(), "libelle": e.title, "montant": str(e.amount), "sens": e.kind,
         "categorie": e.category.label if e.category else None, "compte": e.account.name if e.account else None}
        for e in Entry.objects.filter(created_by=user).order_by("day")[:2000]
    ]


def _documents(user):
    from documents.models import Document

    return [
        {"titre": d.title, "categorie": d.category.name, "depose_le": d.created_at.isoformat()}
        for d in Document.objects.filter(owner=user).select_related("category")[:2000]
    ]


def _availabilities(user):
    from plannings.models import Availability

    return [
        {"campagne": a.campaign.label, "semaine": a.week, "jour": a.weekday, "creneau": a.slot,
         "soir": a.evening, "maj": a.updated_at.isoformat()}
        for a in Availability.objects.filter(member=user).select_related("campaign")[:2000]
    ]


def _assignments(user):
    from chores.models import Assignment

    return [
        {"campagne": a.campaign.label, "semaine": a.week, "jour": a.day.label if a.day else None,
         "tache": a.task.label if a.task else None, "etat": a.status}
        for a in Assignment.objects.filter(member=user).select_related("campaign", "day", "task")[:2000]
    ]


def _messages(user):
    from mail.models import Recipient

    return [
        {"objet": r.broadcast.subject, "lu_le": r.read_at.isoformat() if r.read_at else None}
        for r in Recipient.objects.filter(user=user).select_related("broadcast")[:2000]
    ]


def _logins(user):
    return [
        {"le": la.at.isoformat(), "ip": la.ip, "succes": la.success, "motif": la.reason}
        for la in user.login_attempts.order_by("-at")[:200]
    ]
